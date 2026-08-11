"""
===========================================================
TechDocAI
Servicio de exportación de informes
Entrega 4.4
===========================================================
"""

from pathlib import Path
from datetime import datetime
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import (
    getSampleStyleSheet,
    ParagraphStyle,
)
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)

from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

from app.controllers.document_controller import DocumentController
from app.services.history_service import HistoryService

from config import (
    APP_NAME,
    APP_VERSION,
    AI_PROVIDER,
    MODEL_NAME,
)


class ExportService:
    """
    Servicio encargado de generar informes Markdown,
    PDF y DOCX a partir del documento actualmente procesado.
    """

    # ======================================================
    # DIRECTORIO DE EXPORTACIÓN
    # ======================================================

    BASE_DIR = Path(__file__).resolve().parents[2]

    EXPORT_DIR = BASE_DIR / "exports"

    # ======================================================
    # VALIDACIÓN
    # ======================================================

    @classmethod
    def _obtener_documento(cls):

        documento = DocumentController.documento_actual

        if documento is None:

            raise Exception(
                "No existe un documento cargado."
            )

        return documento

    @classmethod
    def _obtener_analisis(cls, documento):

        analisis = getattr(
            documento,
            "analisis",
            None
        )

        if not analisis:

            raise Exception(
                "Debe analizar el documento antes de exportarlo."
            )

        return analisis

    # ======================================================
    # INFORMACIÓN COMÚN DEL INFORME
    # ======================================================

    @staticmethod
    def _fecha_actual():

        return datetime.now().strftime(
            "%d/%m/%Y %H:%M:%S"
        )

    @staticmethod
    def _fecha_archivo():

        return datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )

    @staticmethod
    def _metadatos_sistema():

        return {
            "Aplicación": APP_NAME,
            "Versión": APP_VERSION,
            "Proveedor IA": AI_PROVIDER,
            "Modelo IA": MODEL_NAME,
        }

    # ======================================================
    # NOMBRE DE ARCHIVO
    # ======================================================

    @classmethod
    def _nombre_archivo(cls, documento):

        nombre = Path(
            documento.nombre
        ).stem

        return nombre

    # ======================================================
    # EXPORTACIÓN MARKDOWN
    # ======================================================

    @classmethod
    def exportar_markdown(cls):

        documento = cls._obtener_documento()

        analisis = cls._obtener_analisis(
            documento
        )

        cls.EXPORT_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        nombre = cls._nombre_archivo(
            documento
        )

        ruta = (
            cls.EXPORT_DIR
            / f"{nombre}_Informe.md"
        )

        fecha = cls._fecha_actual()

        contenido = f"""# {APP_NAME}

## INFORME TÉCNICO GENERADO MEDIANTE INTELIGENCIA ARTIFICIAL

**Documento:** {documento.nombre}

**Fecha de generación:** {fecha}

**Aplicación:** {APP_NAME}

**Versión:** {APP_VERSION}

**Proveedor IA:** {AI_PROVIDER}

**Modelo IA:** {MODEL_NAME}

---

# 1. Información del documento

| Parámetro | Valor |
|---|---|
| Nombre | {documento.nombre} |
| Páginas | {documento.paginas} |
| Tamaño | {documento.tamano_mb:.2f} MB |
| Caracteres | {documento.caracteres:,} |
| Palabras | {documento.palabras:,} |
| Líneas | {documento.lineas:,} |
| Chunks procesados | {documento.chunks} |

---

# 2. Análisis generado por IA

{analisis}

---

## Información del sistema

| Parámetro | Valor |
|---|---|
| Aplicación | {APP_NAME} |
| Versión | {APP_VERSION} |
| Proveedor IA | {AI_PROVIDER} |
| Modelo | {MODEL_NAME} |
| Fecha de generación | {fecha} |

---

**Informe generado automáticamente por TechDocAI.**
"""

        ruta.write_text(
            contenido,
            encoding="utf-8"
        )

        HistoryService.registrar_exportacion(
            documento=documento,
            ruta=ruta,
            formato="Markdown"
        )

        return ruta

    # ======================================================
    # LIMPIEZA DE MARKDOWN PARA REPORTLAB
    # ======================================================

    @staticmethod
    def _formatear_inline(texto):

        texto = escape(
            str(texto)
        )

        # Negrita
        while "**" in texto:

            partes = texto.split(
                "**",
                2
            )

            if len(partes) < 3:
                break

            texto = (
                partes[0]
                + "<b>"
                + partes[1]
                + "</b>"
                + partes[2]
            )

        # Cursiva sencilla
        if (
            texto.startswith("*")
            and texto.endswith("*")
            and not texto.startswith("**")
        ):

            contenido = texto[1:-1]

            if contenido.strip():

                texto = (
                    "<i>"
                    + contenido
                    + "</i>"
                )

        return texto

    # ======================================================
    # ESTILOS PDF
    # ======================================================

    @classmethod
    def _crear_estilos(cls):

        estilos = getSampleStyleSheet()

        estilo_portada = ParagraphStyle(
            "PortadaTechDocAI",
            parent=estilos["Title"],
            fontName="Helvetica-Bold",
            fontSize=25,
            leading=30,
            alignment=TA_CENTER,
            spaceAfter=20,
        )

        estilo_subtitulo = ParagraphStyle(
            "SubtituloTechDocAI",
            parent=estilos["Normal"],
            fontName="Helvetica-Bold",
            fontSize=13,
            leading=18,
            alignment=TA_CENTER,
            spaceAfter=20,
        )

        estilo_documento = ParagraphStyle(
            "DocumentoTechDocAI",
            parent=estilos["Normal"],
            fontName="Helvetica",
            fontSize=12,
            leading=18,
            alignment=TA_CENTER,
            spaceAfter=12,
        )

        estilo_metadato = ParagraphStyle(
            "MetadatoTechDocAI",
            parent=estilos["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=13,
            alignment=TA_CENTER,
            textColor=colors.grey,
            spaceAfter=5,
        )

        estilo_h1 = ParagraphStyle(
            "H1TechDocAI",
            parent=estilos["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=20,
            spaceBefore=10,
            spaceAfter=12,
        )

        estilo_h2 = ParagraphStyle(
            "H2TechDocAI",
            parent=estilos["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=13,
            leading=17,
            spaceBefore=8,
            spaceAfter=8,
        )

        estilo_h3 = ParagraphStyle(
            "H3TechDocAI",
            parent=estilos["Heading3"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=15,
            spaceBefore=6,
            spaceAfter=6,
        )

        estilo_cuerpo = ParagraphStyle(
            "CuerpoTechDocAI",
            parent=estilos["BodyText"],
            fontName="Helvetica",
            fontSize=10,
            leading=15,
            alignment=TA_JUSTIFY,
            spaceAfter=8,
        )

        estilo_lista = ParagraphStyle(
            "ListaTechDocAI",
            parent=estilo_cuerpo,
            leftIndent=15,
            firstLineIndent=-8,
            spaceAfter=5,
        )

        estilo_tabla = ParagraphStyle(
            "TablaTechDocAI",
            parent=estilos["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
        )

        estilo_tabla_negrita = ParagraphStyle(
            "TablaNegritaTechDocAI",
            parent=estilo_tabla,
            fontName="Helvetica-Bold",
        )

        estilo_pie = ParagraphStyle(
            "PieTechDocAI",
            parent=estilos["Normal"],
            fontName="Helvetica",
            fontSize=8,
            textColor=colors.grey,
            alignment=TA_CENTER,
        )

        return {
            "portada": estilo_portada,
            "subtitulo": estilo_subtitulo,
            "documento": estilo_documento,
            "metadato": estilo_metadato,
            "h1": estilo_h1,
            "h2": estilo_h2,
            "h3": estilo_h3,
            "cuerpo": estilo_cuerpo,
            "lista": estilo_lista,
            "tabla": estilo_tabla,
            "tabla_negrita": estilo_tabla_negrita,
            "pie": estilo_pie,
        }

    # ======================================================
    # PIE DE PÁGINA PDF
    # ======================================================

    @staticmethod
    def _pie_pagina(canvas, documento):

        canvas.saveState()

        ancho, alto = A4

        canvas.setFont(
            "Helvetica",
            8
        )

        canvas.setFillColor(
            colors.grey
        )

        canvas.drawCentredString(
            ancho / 2,
            1.0 * cm,
            f"{APP_NAME} v{APP_VERSION} | "
            "Informe generado automáticamente"
        )

        canvas.drawRightString(
            ancho - 1.5 * cm,
            1.0 * cm,
            f"Página {canvas.getPageNumber()}"
        )

        canvas.restoreState()

    # ======================================================
    # METADATOS PDF
    # ======================================================

    @staticmethod
    def _metadatos_pdf(canvas, documento):

        canvas.setTitle(
            f"Informe Técnico - {documento.nombre}"
        )

        canvas.setAuthor(
            APP_NAME
        )

        canvas.setSubject(
            "Informe técnico generado mediante Inteligencia Artificial"
        )

        canvas.setKeywords(
            f"{APP_NAME}, {AI_PROVIDER}, "
            f"{MODEL_NAME}, informe técnico"
        )

    # ======================================================
    # TABLA DE INFORMACIÓN
    # ======================================================

    @classmethod
    def _crear_tabla_documento(
        cls,
        documento,
        estilos
    ):

        datos = [

            [
                Paragraph(
                    "Parámetro",
                    estilos["tabla_negrita"]
                ),
                Paragraph(
                    "Valor",
                    estilos["tabla_negrita"]
                ),
            ],

            [
                Paragraph(
                    "Nombre",
                    estilos["tabla_negrita"]
                ),
                Paragraph(
                    cls._formatear_inline(
                        documento.nombre
                    ),
                    estilos["tabla"]
                ),
            ],

            [
                Paragraph(
                    "Páginas",
                    estilos["tabla_negrita"]
                ),
                Paragraph(
                    str(documento.paginas),
                    estilos["tabla"]
                ),
            ],

            [
                Paragraph(
                    "Tamaño",
                    estilos["tabla_negrita"]
                ),
                Paragraph(
                    f"{documento.tamano_mb:.2f} MB",
                    estilos["tabla"]
                ),
            ],

            [
                Paragraph(
                    "Caracteres",
                    estilos["tabla_negrita"]
                ),
                Paragraph(
                    f"{documento.caracteres:,}",
                    estilos["tabla"]
                ),
            ],

            [
                Paragraph(
                    "Palabras",
                    estilos["tabla_negrita"]
                ),
                Paragraph(
                    f"{documento.palabras:,}",
                    estilos["tabla"]
                ),
            ],

            [
                Paragraph(
                    "Líneas",
                    estilos["tabla_negrita"]
                ),
                Paragraph(
                    f"{documento.lineas:,}",
                    estilos["tabla"]
                ),
            ],

            [
                Paragraph(
                    "Chunks procesados",
                    estilos["tabla_negrita"]
                ),
                Paragraph(
                    str(documento.chunks),
                    estilos["tabla"]
                ),
            ],

            [
                Paragraph(
                    "Proveedor IA",
                    estilos["tabla_negrita"]
                ),
                Paragraph(
                    AI_PROVIDER,
                    estilos["tabla"]
                ),
            ],

            [
                Paragraph(
                    "Modelo IA",
                    estilos["tabla_negrita"]
                ),
                Paragraph(
                    MODEL_NAME,
                    estilos["tabla"]
                ),
            ],

            [
                Paragraph(
                    "Versión TechDocAI",
                    estilos["tabla_negrita"]
                ),
                Paragraph(
                    APP_VERSION,
                    estilos["tabla"]
                ),
            ],
        ]

        tabla = Table(
            datos,
            colWidths=[
                5.0 * cm,
                11.5 * cm,
            ],
            repeatRows=1,
        )

        tabla.setStyle(
            TableStyle(
                [

                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            "#285B86"
                        ),
                    ),

                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),

                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey,
                    ),

                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "MIDDLE",
                    ),

                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),

                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),

                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),

                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                ]
            )
        )

        return tabla

    # ======================================================
    # ELEMENTOS DEL ANÁLISIS PARA PDF
    # ======================================================

    @classmethod
    def _analisis_a_elementos(
        cls,
        analisis,
        estilos
    ):

        elementos = []

        lineas = str(
            analisis
        ).splitlines()

        for linea in lineas:

            linea = linea.strip()

            # Línea vacía
            if not linea:

                elementos.append(
                    Spacer(
                        1,
                        0.12 * cm
                    )
                )

                continue

            # H1
            if (
                linea.startswith("# ")
                and not linea.startswith("## ")
            ):

                texto = linea[2:].strip()

                elementos.append(
                    Paragraph(
                        cls._formatear_inline(
                            texto
                        ),
                        estilos["h1"]
                    )
                )

                continue

            # H2
            if (
                linea.startswith("## ")
                and not linea.startswith("### ")
            ):

                texto = linea[3:].strip()

                elementos.append(
                    Paragraph(
                        cls._formatear_inline(
                            texto
                        ),
                        estilos["h2"]
                    )
                )

                continue

            # H3
            if linea.startswith("### "):

                texto = linea[4:].strip()

                elementos.append(
                    Paragraph(
                        cls._formatear_inline(
                            texto
                        ),
                        estilos["h3"]
                    )
                )

                continue

            # Lista con guion
            if (
                linea.startswith("- ")
                or linea.startswith("* ")
            ):

                texto = linea[2:].strip()

                elementos.append(
                    Paragraph(
                        "• "
                        + cls._formatear_inline(
                            texto
                        ),
                        estilos["lista"]
                    )
                )

                continue

            # Lista numerada
            es_lista_numerada = (
                len(linea) >= 3
                and linea[0].isdigit()
                and linea[1:3] == ". "
            )

            if es_lista_numerada:

                elementos.append(
                    Paragraph(
                        cls._formatear_inline(
                            linea
                        ),
                        estilos["lista"]
                    )
                )

                continue

            # Separador Markdown
            if linea in (
                "---",
                "***",
                "___"
            ):

                elementos.append(
                    Spacer(
                        1,
                        0.2 * cm
                    )
                )

                continue

            # Párrafo normal
            texto = cls._formatear_inline(
                linea
            )

            elementos.append(
                Paragraph(
                    texto,
                    estilos["cuerpo"]
                )
            )

        return elementos

    # ======================================================
    # EXPORTAR PDF
    # ======================================================

    @classmethod
    def exportar_pdf(cls):

        documento = cls._obtener_documento()

        analisis = cls._obtener_analisis(
            documento
        )

        cls.EXPORT_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        nombre = cls._nombre_archivo(
            documento
        )

        ruta = (
            cls.EXPORT_DIR
            / f"{nombre}_Informe.pdf"
        )

        estilos = cls._crear_estilos()

        doc = SimpleDocTemplate(
            str(ruta),
            pagesize=A4,
            rightMargin=1.8 * cm,
            leftMargin=1.8 * cm,
            topMargin=1.8 * cm,
            bottomMargin=1.8 * cm,
            title=f"Informe Técnico - {documento.nombre}",
            author=APP_NAME,
        )

        elementos = []

        # ==================================================
        # PORTADA
        # ==================================================

        elementos.append(
            Spacer(
                1,
                3.5 * cm
            )
        )

        elementos.append(
            Paragraph(
                APP_NAME.upper(),
                estilos["portada"]
            )
        )

        elementos.append(
            Paragraph(
                "INFORME TÉCNICO GENERADO "
                "MEDIANTE INTELIGENCIA ARTIFICIAL",
                estilos["subtitulo"]
            )
        )

        elementos.append(
            Spacer(
                1,
                1.2 * cm
            )
        )

        elementos.append(
            Paragraph(
                cls._formatear_inline(
                    documento.nombre
                ),
                estilos["documento"]
            )
        )

        fecha = cls._fecha_actual()

        elementos.append(
            Paragraph(
                f"Fecha de generación: {fecha}",
                estilos["documento"]
            )
        )

        elementos.append(
            Paragraph(
                f"Versión: {APP_VERSION}",
                estilos["metadato"]
            )
        )

        elementos.append(
            Paragraph(
                f"Proveedor IA: {AI_PROVIDER}",
                estilos["metadato"]
            )
        )

        elementos.append(
            Paragraph(
                f"Modelo: {MODEL_NAME}",
                estilos["metadato"]
            )
        )

        elementos.append(
            Spacer(
                1,
                5.0 * cm
            )
        )

        elementos.append(
            Paragraph(
                "Sistema Inteligente de Análisis "
                "de Documentos Técnicos",
                estilos["documento"]
            )
        )

        # ==================================================
        # INFORMACIÓN DEL DOCUMENTO
        # ==================================================

        elementos.append(
            PageBreak()
        )

        elementos.append(
            Paragraph(
                "1. Información del documento",
                estilos["h1"]
            )
        )

        elementos.append(
            Spacer(
                1,
                0.2 * cm
            )
        )

        elementos.append(
            cls._crear_tabla_documento(
                documento,
                estilos
            )
        )

        elementos.append(
            Spacer(
                1,
                0.7 * cm
            )
        )

        # ==================================================
        # ANÁLISIS
        # ==================================================

        elementos.append(
            Paragraph(
                "2. Análisis generado por IA",
                estilos["h1"]
            )
        )

        elementos.extend(
            cls._analisis_a_elementos(
                analisis,
                estilos
            )
        )

        # ==================================================
        # CONSTRUCCIÓN PDF
        # ==================================================

        def configurar_pdf(
            canvas,
            doc_pdf
        ):

            cls._metadatos_pdf(
                canvas,
                documento
            )

            cls._pie_pagina(
                canvas,
                documento
            )

        doc.build(
            elementos,
            onFirstPage=configurar_pdf,
            onLaterPages=configurar_pdf,
        )

        HistoryService.registrar_exportacion(
            documento=documento,
            ruta=ruta,
            formato="PDF"
        )

        return ruta

    # ======================================================
    # FORMATEAR PÁRRAFO DOCX
    # ======================================================

    @staticmethod
    def _formatear_parrafo_docx(
        parrafo,
        texto,
        justificado=True
    ):

        if justificado:

            parrafo.alignment = (
                WD_ALIGN_PARAGRAPH.JUSTIFY
            )

        run = parrafo.add_run(
            texto
        )

        run.font.name = "Arial"
        run.font.size = Pt(10)

        return run

    # ======================================================
    # CONFIGURAR ENCABEZADO DOCX
    # ======================================================

    @classmethod
    def _configurar_encabezado_docx(
        cls,
        seccion
    ):

        encabezado = seccion.header

        parrafo = encabezado.paragraphs[0]

        parrafo.alignment = (
            WD_ALIGN_PARAGRAPH.RIGHT
        )

        run = parrafo.add_run(
            f"{APP_NAME} v{APP_VERSION}"
        )

        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.italic = True

    # ======================================================
    # CONFIGURAR PIE DOCX
    # ======================================================

    @classmethod
    def _configurar_pie_docx(
        cls,
        seccion
    ):

        pie = seccion.footer

        parrafo = pie.paragraphs[0]

        parrafo.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = parrafo.add_run(
            f"Informe generado automáticamente por "
            f"{APP_NAME} | v{APP_VERSION}"
        )

        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.italic = True

    # ======================================================
    # EXPORTAR DOCX
    # ======================================================

    @classmethod
    def exportar_docx(cls):

        documento = cls._obtener_documento()

        analisis = cls._obtener_analisis(
            documento
        )

        cls.EXPORT_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        nombre = cls._nombre_archivo(
            documento
        )

        ruta = (
            cls.EXPORT_DIR
            / f"{nombre}_Informe.docx"
        )

        word = WordDocument()

        # ==================================================
        # PROPIEDADES INTERNAS DOCX
        # ==================================================

        propiedades = word.core_properties

        propiedades.title = (
            f"Informe Técnico - {documento.nombre}"
        )

        propiedades.subject = (
            "Informe técnico generado mediante "
            "Inteligencia Artificial"
        )

        propiedades.author = APP_NAME

        propiedades.keywords = (
            f"{APP_NAME}, {AI_PROVIDER}, "
            f"{MODEL_NAME}, informe técnico"
        )

        propiedades.comments = (
            f"Generado por {APP_NAME} "
            f"v{APP_VERSION}"
        )

        # ==================================================
        # CONFIGURACIÓN DE PÁGINA
        # ==================================================

        seccion = word.sections[0]

        seccion.top_margin = Cm(2)
        seccion.bottom_margin = Cm(2)
        seccion.left_margin = Cm(2)
        seccion.right_margin = Cm(2)

        cls._configurar_encabezado_docx(
            seccion
        )

        cls._configurar_pie_docx(
            seccion
        )

        # ==================================================
        # ESTILOS
        # ==================================================

        estilo_normal = word.styles["Normal"]

        estilo_normal.font.name = "Arial"
        estilo_normal.font.size = Pt(10)

        # ==================================================
        # PORTADA
        # ==================================================

        word.add_paragraph()
        word.add_paragraph()
        word.add_paragraph()

        titulo = word.add_paragraph()

        titulo.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = titulo.add_run(
            APP_NAME.upper()
        )

        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(24)

        subtitulo = word.add_paragraph()

        subtitulo.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = subtitulo.add_run(
            "INFORME TÉCNICO GENERADO "
            "MEDIANTE INTELIGENCIA ARTIFICIAL"
        )

        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(13)

        word.add_paragraph()

        nombre_doc = word.add_paragraph()

        nombre_doc.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = nombre_doc.add_run(
            documento.nombre
        )

        run.font.name = "Arial"
        run.font.size = Pt(12)

        fecha = cls._fecha_actual()

        fecha_parrafo = word.add_paragraph()

        fecha_parrafo.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = fecha_parrafo.add_run(
            f"Fecha de generación: {fecha}"
        )

        run.font.name = "Arial"
        run.font.size = Pt(10)

        word.add_paragraph()

        version_parrafo = word.add_paragraph()

        version_parrafo.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = version_parrafo.add_run(
            f"TechDocAI v{APP_VERSION}"
        )

        run.font.name = "Arial"
        run.font.size = Pt(9)

        proveedor_parrafo = word.add_paragraph()

        proveedor_parrafo.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = proveedor_parrafo.add_run(
            f"Proveedor IA: {AI_PROVIDER}"
        )

        run.font.name = "Arial"
        run.font.size = Pt(9)

        modelo_parrafo = word.add_paragraph()

        modelo_parrafo.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = modelo_parrafo.add_run(
            f"Modelo: {MODEL_NAME}"
        )

        run.font.name = "Arial"
        run.font.size = Pt(9)

        word.add_page_break()

        # ==================================================
        # INFORMACIÓN DEL DOCUMENTO
        # ==================================================

        word.add_heading(
            "1. Información del documento",
            level=1
        )

        datos = [

            (
                "Nombre",
                documento.nombre
            ),

            (
                "Páginas",
                str(documento.paginas)
            ),

            (
                "Tamaño",
                f"{documento.tamano_mb:.2f} MB"
            ),

            (
                "Caracteres",
                f"{documento.caracteres:,}"
            ),

            (
                "Palabras",
                f"{documento.palabras:,}"
            ),

            (
                "Líneas",
                f"{documento.lineas:,}"
            ),

            (
                "Chunks procesados",
                str(documento.chunks)
            ),

            (
                "Proveedor IA",
                AI_PROVIDER
            ),

            (
                "Modelo IA",
                MODEL_NAME
            ),

            (
                "Versión TechDocAI",
                APP_VERSION
            ),

            (
                "Fecha de generación",
                fecha
            ),
        ]

        tabla = word.add_table(
            rows=1,
            cols=2
        )

        tabla.style = "Table Grid"

        encabezado = tabla.rows[0].cells

        encabezado[0].text = "Parámetro"
        encabezado[1].text = "Valor"

        for celda in encabezado:

            for parrafo in celda.paragraphs:

                for run in parrafo.runs:

                    run.bold = True
                    run.font.name = "Arial"
                    run.font.size = Pt(9)

        for parametro, valor in datos:

            celdas = tabla.add_row().cells

            celdas[0].text = parametro
            celdas[1].text = valor

            for celda in celdas:

                for parrafo in celda.paragraphs:

                    for run in parrafo.runs:

                        run.font.name = "Arial"
                        run.font.size = Pt(9)

        word.add_paragraph()

        # ==================================================
        # ANÁLISIS IA
        # ==================================================

        word.add_heading(
            "2. Análisis generado por IA",
            level=1
        )

        numero_lista = 0

        lineas = str(
            analisis
        ).splitlines()

        for linea in lineas:

            linea = linea.strip()

            # ----------------------------------------------
            # LÍNEA VACÍA
            # ----------------------------------------------

            if not linea:

                numero_lista = 0

                word.add_paragraph()

                continue

            # ----------------------------------------------
            # H1
            # ----------------------------------------------

            if (
                linea.startswith("# ")
                and not linea.startswith("## ")
            ):

                numero_lista = 0

                texto = linea[2:].strip()

                word.add_heading(
                    texto.replace("**", ""),
                    level=1
                )

                continue

            # ----------------------------------------------
            # H2
            # ----------------------------------------------

            if (
                linea.startswith("## ")
                and not linea.startswith("### ")
            ):

                numero_lista = 0

                texto = linea[3:].strip()

                word.add_heading(
                    texto.replace("**", ""),
                    level=2
                )

                continue

            # ----------------------------------------------
            # H3
            # ----------------------------------------------

            if linea.startswith("### "):

                numero_lista = 0

                texto = linea[4:].strip()

                word.add_heading(
                    texto.replace("**", ""),
                    level=3
                )

                continue

            # ----------------------------------------------
            # LISTA CON VIÑETAS
            # ----------------------------------------------

            if (
                linea.startswith("- ")
                or linea.startswith("* ")
            ):

                numero_lista = 0

                texto = linea[2:].strip()

                parrafo = word.add_paragraph()

                parrafo.style = (
                    word.styles["List Bullet"]
                )

                cls._formatear_parrafo_docx(
                    parrafo,
                    texto.replace("**", "")
                )

                continue

            # ----------------------------------------------
            # LISTA NUMERADA
            # ----------------------------------------------

            es_numerada = (
                len(linea) >= 3
                and linea[0].isdigit()
                and linea[1:3] == ". "
            )

            if es_numerada:

                numero_lista += 1

                texto = linea[3:].strip()

                parrafo = word.add_paragraph()

                cls._formatear_parrafo_docx(
                    parrafo,
                    f"{numero_lista}. "
                    f"{texto.replace('**', '')}"
                )

                continue

            # ----------------------------------------------
            # PÁRRAFO NORMAL
            # ----------------------------------------------

            numero_lista = 0

            parrafo = word.add_paragraph()

            cls._formatear_parrafo_docx(
                parrafo,
                linea.replace("**", "")
            )

        # ==================================================
        # PIE FINAL
        # ==================================================

        word.add_paragraph()

        pie = word.add_paragraph()

        pie.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = pie.add_run(
            f"Informe generado automáticamente por "
            f"{APP_NAME}."
        )

        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.italic = True

        # ==================================================
        # GUARDAR
        # ==================================================

        word.save(
            str(ruta)
        )

        HistoryService.registrar_exportacion(
            documento=documento,
            ruta=ruta,
            formato="DOCX"
        )

        return ruta